#!/usr/bin/env python3
"""CHEEZE 캡스톤 기말 발표 PPT 생성기

슬라이드 구성 (10슬라이드):
  1. 표지
  2. 중간발표 피드백 반영
  3. 보완 1: 통합 로그인 시스템
  4. 보완 2: 보안 취약점 분석 및 강화
  5. 인프라 정비 및 최적화
  6. LIVE DEMO
  7. 기대효과
  8. 개발 일정
  9. 향후 보완 과제
  10. 감사합니다

폰트: 맑은 고딕 (Malgun Gothic) + Calibri — Windows 기본 내장 폰트만 사용
슬라이드 크기: 16:9 와이드스크린 (13.333 x 7.5 인치)
"""
import sys, os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', '..', '.vendor'))

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# ── 색상 팔레트 ──
FONT_KR = '맑은 고딕'
FONT_EN = 'Calibri'
DARK    = RGBColor(0x1A, 0x1A, 0x2E)
CARD    = RGBColor(0x25, 0x25, 0x3A)
BLUE    = RGBColor(0x4F, 0x7F, 0xFF)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
GRAY    = RGBColor(0x6B, 0x72, 0x80)
GREEN   = RGBColor(0x10, 0xB9, 0x81)
ORANGE  = RGBColor(0xF5, 0x9E, 0x0B)
LIGHT_BLUE = RGBColor(0xA0, 0xB8, 0xFF)

# ── 프레젠테이션 초기화 ──
prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)


# ────────────────────────────────────────────────
# 헬퍼 함수
# ────────────────────────────────────────────────

def blank_slide():
    return prs.slides.add_slide(prs.slide_layouts[6])


def set_bg(slide, color=DARK):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text(slide, left, top, width, height,
             text, size=18, color=WHITE, bold=False,
             align=PP_ALIGN.LEFT, font=FONT_KR):
    txb = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font
    p.alignment = align
    return txb


def add_para(tf, text, size=16, color=WHITE, bold=False,
             space_before=Pt(6), font=FONT_KR):
    p = tf.add_paragraph()
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font
    p.space_before = space_before
    return p


def add_rect(slide, left, top, width, height, fill_color, line=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def add_sharp_rect(slide, left, top, width, height, fill_color, line=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def slide_title(slide, text, sub=None):
    """슬라이드 상단 제목 + 선"""
    add_text(slide, 0.8, 0.35, 11.7, 0.7, text,
             size=32, color=WHITE, bold=True)
    # 밑줄
    add_sharp_rect(slide, 0.8, 1.05, 11.7, 0.04, BLUE)
    if sub:
        add_text(slide, 0.8, 1.15, 11.7, 0.45, sub,
                 size=15, color=GRAY)


# ────────────────────────────────────────────────
# Slide 1: 표지
# ────────────────────────────────────────────────
slide = blank_slide()
set_bg(slide, DARK)

# 배경 장식 직사각형
add_sharp_rect(slide, 0, 0, 0.5, 7.5, BLUE)
add_sharp_rect(slide, 12.833, 0, 0.5, 7.5, BLUE)

add_text(slide, 1, 1.2, 11.333, 1.2,
         '클라우드 기반 통합 서비스 플랫폼',
         size=40, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
add_text(slide, 1, 2.7, 11.333, 0.7,
         '캡스톤디자인 기말 발표',
         size=24, color=LIGHT_BLUE, align=PP_ALIGN.CENTER)
add_sharp_rect(slide, 4.5, 3.6, 4.333, 0.05, BLUE)
add_text(slide, 1, 3.9, 11.333, 0.6,
         '살려조  |  팀장 정성현',
         size=18, color=GRAY, align=PP_ALIGN.CENTER)
add_text(slide, 1, 4.7, 11.333, 0.6,
         '2026. 06. 11',
         size=18, color=GRAY, align=PP_ALIGN.CENTER, font=FONT_EN)


# ────────────────────────────────────────────────
# Slide 2: 중간발표 피드백 반영
# ────────────────────────────────────────────────
slide = blank_slide()
set_bg(slide, DARK)
slide_title(slide, '중간발표 피드백 반영')

# 왼쪽: 완료 항목
add_rect(slide, 0.7, 1.7, 5.9, 5.2, CARD)
add_text(slide, 1.0, 1.85, 5.3, 0.45,
         '중간발표 구현 완료 항목',
         size=16, color=BLUE, bold=True)

done = [
    '통합 포탈 (메인 페이지, 서비스 검색)',
    '클라우드 서비스 연동 (Nextcloud, Paperless, ArchiveBox)',
    '온디맨드 서버 제어 (WOL, 자동 종료)',
    '관리자 대시보드 (서비스 상태, 감사 로그,\n   절전 관리, 모니터링)',
    '학습 플랫폼 (AWS SAA CBT)',
]
for i, item in enumerate(done):
    add_text(slide, 1.15, 2.45 + i * 0.75, 0.35, 0.45,
             '✓', size=14, color=GREEN, font=FONT_EN)
    add_text(slide, 1.55, 2.45 + i * 0.75, 4.8, 0.65,
             item, size=15, color=WHITE)

# 오른쪽: 보완 과제 (강조 카드 2개)
add_text(slide, 7.0, 1.85, 5.5, 0.45,
         '보완 과제',
         size=16, color=ORANGE, bold=True)

tasks = [
    ('01', '통합 로그인 시스템',
     '서비스마다 별도 계정 → 단일 포탈 통합 인증'),
    ('02', '보안 취약점 분석 및 강화',
     '자체 호스팅 인프라 보안 체계 구축'),
]
for i, (num, title, desc) in enumerate(tasks):
    y = 2.5 + i * 2.1
    add_rect(slide, 7.0, y, 5.7, 1.8, RGBColor(0x2A, 0x1A, 0x0A))
    add_sharp_rect(slide, 7.0, y, 0.08, 1.8, ORANGE)
    add_text(slide, 7.25, y + 0.15, 1.0, 0.45,
             num, size=28, color=ORANGE, bold=True, font=FONT_EN)
    add_text(slide, 7.25, y + 0.65, 5.0, 0.45,
             title, size=18, color=WHITE, bold=True)
    add_text(slide, 7.25, y + 1.15, 5.0, 0.45,
             desc, size=13, color=GRAY)


# ────────────────────────────────────────────────
# Slide 3: 보완 1 — 통합 로그인 시스템
# ────────────────────────────────────────────────
slide = blank_slide()
set_bg(slide, DARK)
slide_title(slide, '보완 1: 통합 로그인 시스템',
            sub='서비스별 분산 계정 → 포탈 단일 인증으로 통합')

# Before 카드
add_rect(slide, 0.7, 1.7, 5.5, 4.0, RGBColor(0x3A, 0x15, 0x15))
add_sharp_rect(slide, 0.7, 1.7, 5.5, 0.05, ORANGE)
add_text(slide, 1.0, 1.85, 4.8, 0.45,
         'BEFORE', size=16, color=ORANGE, bold=True, font=FONT_EN)

before_items = [
    'Nextcloud 로그인',
    'Paperless 로그인',
    'ArchiveBox 로그인',
    '→ 3번 별도 로그인 필요',
]
colors_b = [GRAY, GRAY, GRAY, ORANGE]
for i, (item, c) in enumerate(zip(before_items, colors_b)):
    add_text(slide, 1.3, 2.45 + i * 0.6, 4.5, 0.5,
             item, size=15, color=c, bold=(c == ORANGE))

# After 카드
add_rect(slide, 6.8, 1.7, 5.8, 4.0, RGBColor(0x0A, 0x20, 0x30))
add_sharp_rect(slide, 6.8, 1.7, 5.8, 0.05, BLUE)
add_text(slide, 7.1, 1.85, 5.0, 0.45,
         'AFTER', size=16, color=BLUE, bold=True, font=FONT_EN)

after_items = [
    '포탈 통합 로그인 (1번)',
    '역할 기반 접근 (관리자 / 직원)',
    '서비스 대시보드에서 바로 접속',
]
for i, item in enumerate(after_items):
    add_text(slide, 7.4, 2.45 + i * 0.6, 4.8, 0.5,
             item, size=15, color=WHITE)

# 구현 내용
add_rect(slide, 6.8, 4.0, 5.8, 1.3, CARD)
add_text(slide, 7.1, 4.1, 5.2, 0.4,
         '구현 내용', size=14, color=GREEN, bold=True)
add_text(slide, 7.1, 4.55, 5.2, 0.7,
         'JWT 세션 토큰  |  8시간 TTL  |  역할별 UI 분기',
         size=13, color=GRAY)

# 화살표 구분선
add_text(slide, 5.85, 3.3, 1.0, 0.5,
         '→', size=36, color=BLUE, bold=True, font=FONT_EN,
         align=PP_ALIGN.CENTER)

# 향후 계획 배너
add_rect(slide, 0.7, 6.1, 11.9, 0.75, RGBColor(0x0A, 0x14, 0x2A))
add_text(slide, 1.0, 6.2, 11.5, 0.5,
         '향후: OIDC 기반 SSO로 각 서비스 자동 로그인까지 확장 예정',
         size=14, color=LIGHT_BLUE, align=PP_ALIGN.CENTER)


# ────────────────────────────────────────────────
# Slide 4: 보완 2 — 보안 취약점 분석 및 강화
# ────────────────────────────────────────────────
slide = blank_slide()
set_bg(slide, DARK)
slide_title(slide, '보완 2: 보안 취약점 분석 및 강화',
            sub='자체 호스팅 인프라를 위한 4계층 Zero Trust 보안 구조')

sec_layers = [
    ('1계층', 'Cloudflare Tunnel',
     '오리진 IP 비노출  |  DDoS 방어  |  자동 TLS',
     RGBColor(0xF4, 0x81, 0x20), RGBColor(0x2A, 0x1A, 0x08)),
    ('2계층', 'Cloudflare Access',
     '관리자 이메일 OTP 인증  |  Zero Trust  |  JWT',
     BLUE, RGBColor(0x0A, 0x14, 0x2A)),
    ('3계층', 'Token API 인증',
     'SHA-256 해시  |  역할별 권한  |  시간 제한',
     GREEN, RGBColor(0x08, 0x20, 0x18)),
    ('4계층', '감사 로그',
     '모든 제어 행위 JSON Lines 기록  |  IP 추적  |  변조 불가',
     RGBColor(0x8B, 0x5C, 0xF6), RGBColor(0x18, 0x10, 0x2A)),
]
for i, (num, title, desc, accent, bg) in enumerate(sec_layers):
    y = 1.6 + i * 1.3
    add_rect(slide, 0.7, y, 11.9, 1.1, bg)
    add_sharp_rect(slide, 0.7, y, 0.1, 1.1, accent)
    add_text(slide, 1.1, y + 0.08, 1.4, 0.4,
             num, size=15, color=accent, bold=True, font=FONT_EN)
    add_text(slide, 2.3, y + 0.05, 4.5, 0.45,
             title, size=20, color=WHITE, bold=True, font=FONT_EN)
    add_text(slide, 2.3, y + 0.55, 9.5, 0.4,
             desc, size=14, color=GRAY)


# ────────────────────────────────────────────────
# Slide 5: 인프라 정비 및 최적화
# ────────────────────────────────────────────────
slide = blank_slide()
set_bg(slide, DARK)
slide_title(slide, '인프라 정비 및 최적화',
            sub='미사용 컴포넌트 제거 및 UI/코드 전면 리팩토링')

infra_cards = [
    ('Pterodactyl 제거',
     '게임서버 관리 시스템 미사용 1달+\n컨테이너 3개 + 볼륨 3개 정리',
     ORANGE),
    ('E-class 자동화 제거',
     '프론트엔드 + 백엔드 컨테이너\nLXC 환경 정리',
     ORANGE),
    ('UI 리디자인',
     '크롬 스타일 개인 대시보드\n→ 역할 기반 업무 포탈',
     BLUE),
    ('코드 정비',
     '-2,175줄 삭제\n31파일 변경',
     GREEN),
]
for i, (title, desc, accent) in enumerate(infra_cards):
    col = i % 2
    row = i // 2
    x = 0.7 + col * 6.3
    y = 1.7 + row * 2.3
    add_rect(slide, x, y, 5.9, 2.0, CARD)
    add_sharp_rect(slide, x, y, 5.9, 0.07, accent)
    add_text(slide, x + 0.3, y + 0.2, 5.3, 0.45,
             title, size=18, color=WHITE, bold=True)
    add_text(slide, x + 0.3, y + 0.75, 5.3, 1.1,
             desc, size=14, color=GRAY)

# 서비스 설명 리프레이밍 배너
add_rect(slide, 0.7, 6.2, 11.9, 0.65, RGBColor(0x0A, 0x14, 0x2A))
add_text(slide, 1.0, 6.3, 11.5, 0.45,
         '서비스 설명 리프레이밍: 개인용 → 업무용  |  소규모 조직 맞춤 메시지로 전환',
         size=13, color=LIGHT_BLUE, align=PP_ALIGN.CENTER)


# ────────────────────────────────────────────────
# Slide 6: LIVE DEMO
# ────────────────────────────────────────────────
slide = blank_slide()
set_bg(slide, RGBColor(0x0A, 0x0A, 0x1A))

# 배경 장식
add_sharp_rect(slide, 0, 3.55, 13.333, 0.05, BLUE)

add_text(slide, 1, 1.2, 11.333, 1.2,
         'LIVE DEMO',
         size=64, color=BLUE, bold=True, align=PP_ALIGN.CENTER, font=FONT_EN)
add_text(slide, 1, 2.7, 11.333, 0.6,
         '실제 운영 중인 시스템 시연',
         size=24, color=WHITE, align=PP_ALIGN.CENTER)

steps = [
    ('01', '포탈 로그인'),
    ('02', '대시보드'),
    ('03', 'Nextcloud'),
    ('04', 'Paperless'),
    ('05', '관리자 콘솔'),
]
total_w = 11.9
card_w = total_w / len(steps)
for i, (num, label) in enumerate(steps):
    x = 0.7 + i * card_w
    add_rect(slide, x + 0.1, 4.1, card_w - 0.2, 1.8, CARD)
    add_text(slide, x + 0.1, 4.25, card_w - 0.2, 0.5,
             num, size=14, color=BLUE, bold=True, font=FONT_EN,
             align=PP_ALIGN.CENTER)
    add_text(slide, x + 0.1, 4.8, card_w - 0.2, 0.8,
             label, size=15, color=WHITE, align=PP_ALIGN.CENTER)


# ────────────────────────────────────────────────
# Slide 7: 기대효과
# ────────────────────────────────────────────────
slide = blank_slide()
set_bg(slide, DARK)
slide_title(slide, '기대효과')

effects = [
    ('서비스 통합',
     '분산된 서비스를 단일 포탈에서 접근\n사용자 편의성 향상',
     BLUE),
    ('보안 강화',
     'Zero Trust 아키텍처로\n자체 호스팅 보안 취약점 해소',
     GREEN),
    ('운영 효율화',
     '온디맨드 서버 제어로 자원 소모 제거\n관리자 대시보드 중앙 관리',
     ORANGE),
    ('확장 가능성',
     '소규모 조직에 동일 구조 적용 가능\nOIDC SSO 등 표준 프로토콜 확장',
     RGBColor(0x8B, 0x5C, 0xF6)),
]
for i, (title, desc, accent) in enumerate(effects):
    col = i % 2
    row = i // 2
    x = 0.7 + col * 6.3
    y = 1.7 + row * 2.5
    add_rect(slide, x, y, 5.9, 2.2, CARD)
    add_sharp_rect(slide, x, y, 0.1, 2.2, accent)
    add_text(slide, x + 0.4, y + 0.25, 5.1, 0.5,
             title, size=20, color=WHITE, bold=True)
    add_text(slide, x + 0.4, y + 0.85, 5.1, 1.1,
             desc, size=14, color=GRAY)


# ────────────────────────────────────────────────
# Slide 8: 개발 일정 (Gantt 스타일)
# ────────────────────────────────────────────────
slide = blank_slide()
set_bg(slide, DARK)
slide_title(slide, '개발 일정')

# ── 테이블 레이아웃 파라미터 ──
TABLE_LEFT   = 0.7
TABLE_TOP    = 1.6
ROW_H        = 0.78
LABEL_W      = 2.6    # 항목 이름 열 너비

# 월별 열: (3월1-4주, 4월1-5주, 5월1-4주, 6월1-2주) → 총 15주
# 각 주를 0.65인치로 배치
WEEK_W       = 0.65
MONTH_WEEKS  = [4, 5, 4, 2]          # 월별 주 수
MONTH_LABELS = ['3월', '4월', '5월', '6월']
MONTH_COLORS = [
    RGBColor(0x1E, 0x26, 0x3A),
    RGBColor(0x1A, 0x22, 0x36),
    RGBColor(0x1E, 0x26, 0x3A),
    RGBColor(0x1A, 0x22, 0x36),
]

total_weeks = sum(MONTH_WEEKS)   # 15
grid_w = total_weeks * WEEK_W    # 9.75

rows = [
    '주제선정 / 자료조사',
    '세부기능 기획 / 사전구현',
    '구현',
    '피드백 / 수정',
    '테스트 / 보고서 / 발표',
]

# 각 행의 진행 구간: (시작 주 인덱스, 끝 주 인덱스, 계획색, 실제색)
# 주 인덱스: 3월1주=0, 3월4주=3, 4월1주=4, 4월5주=8, 5월1주=9, 5월4주=12, 6월1주=13, 6월2주=14
# (start_plan, end_plan+1, start_actual, end_actual+1)
# plan=GREEN(계획), actual=ORANGE(실제) — 대부분 일치하므로 같은 범위
row_data = [
    (0, 2, 0, 2),     # 주제선정/자료조사: 3월1-2주
    (1, 5, 1, 5),     # 세부기능 기획: 3월2주~4월1주
    (4, 13, 4, 13),   # 구현: 4월1주~5월4주
    (8, 13, 8, 13),   # 피드백/수정: 4월5주~5월4주
    (12, 15, 12, 15), # 테스트/보고서: 5월4주~6월2주
]

# 헤더 배경
add_sharp_rect(slide, TABLE_LEFT, TABLE_TOP, LABEL_W, ROW_H,
               RGBColor(0x1E, 0x28, 0x40))
add_text(slide, TABLE_LEFT + 0.1, TABLE_TOP + 0.2, LABEL_W - 0.2, 0.4,
         '항목', size=12, color=LIGHT_BLUE, bold=True)

# 월별 헤더
week_cursor = 0
for m_idx, (months, weeks, bg) in enumerate(
        zip(MONTH_LABELS, MONTH_WEEKS, MONTH_COLORS)):
    mx = TABLE_LEFT + LABEL_W + week_cursor * WEEK_W
    mw = weeks * WEEK_W
    add_sharp_rect(slide, mx, TABLE_TOP, mw, ROW_H * 0.45, bg)
    add_text(slide, mx, TABLE_TOP + 0.05, mw, 0.35,
             months, size=12, color=LIGHT_BLUE, bold=True, align=PP_ALIGN.CENTER)
    # 주 번호
    for w in range(weeks):
        wx = mx + w * WEEK_W
        add_sharp_rect(slide, wx, TABLE_TOP + ROW_H * 0.45,
                       WEEK_W, ROW_H * 0.55, bg)
        add_text(slide, wx, TABLE_TOP + ROW_H * 0.46,
                 WEEK_W, ROW_H * 0.5,
                 f'{w+1}주', size=9, color=GRAY, align=PP_ALIGN.CENTER)
    week_cursor += weeks

# 행
for r_idx, (row_label, (sp, ep, sa, ea)) in enumerate(zip(rows, row_data)):
    ry = TABLE_TOP + ROW_H + r_idx * ROW_H
    # 항목 셀
    bg_row = CARD if r_idx % 2 == 0 else RGBColor(0x20, 0x20, 0x32)
    add_sharp_rect(slide, TABLE_LEFT, ry, LABEL_W, ROW_H, bg_row)
    add_text(slide, TABLE_LEFT + 0.1, ry + 0.2, LABEL_W - 0.15, 0.45,
             row_label, size=12, color=WHITE)
    # 주 셀 배경
    for w in range(total_weeks):
        wx = TABLE_LEFT + LABEL_W + w * WEEK_W
        add_sharp_rect(slide, wx, ry, WEEK_W, ROW_H, bg_row)
    # 계획 막대 (GREEN, 위쪽)
    if ep > sp:
        bx = TABLE_LEFT + LABEL_W + sp * WEEK_W
        bw = (ep - sp) * WEEK_W
        add_sharp_rect(slide, bx + 0.05, ry + 0.1,
                       bw - 0.1, ROW_H * 0.38, GREEN)
    # 실제 막대 (ORANGE, 아래쪽)
    if ea > sa:
        bx = TABLE_LEFT + LABEL_W + sa * WEEK_W
        bw = (ea - sa) * WEEK_W
        add_sharp_rect(slide, bx + 0.05, ry + ROW_H * 0.52,
                       bw - 0.1, ROW_H * 0.38, ORANGE)

# 마일스톤 표시 — 중간발표 (4월4주 = 주 인덱스 7 끝)
mid_x = TABLE_LEFT + LABEL_W + 8 * WEEK_W
add_sharp_rect(slide, mid_x - 0.02, TABLE_TOP + ROW_H,
               0.04, ROW_H * len(rows), RGBColor(0xFF, 0xFF, 0x00))
add_text(slide, mid_x - 0.5, TABLE_TOP + ROW_H * 0.5,
         1.2, 0.35, '중간발표', size=10,
         color=RGBColor(0xFF, 0xFF, 0x00), align=PP_ALIGN.CENTER)

# 기말발표 (6월2주 = 주 인덱스 14 끝)
fin_x = TABLE_LEFT + LABEL_W + 15 * WEEK_W
add_sharp_rect(slide, fin_x - 0.04, TABLE_TOP + ROW_H,
               0.04, ROW_H * len(rows), RGBColor(0xFF, 0x60, 0x60))
add_text(slide, fin_x - 0.85, TABLE_TOP + ROW_H * 0.5,
         1.2, 0.35, '기말발표', size=10,
         color=RGBColor(0xFF, 0x60, 0x60), align=PP_ALIGN.CENTER)

# 범례
legend_y = TABLE_TOP + ROW_H * (len(rows) + 1) + 0.1
add_sharp_rect(slide, 0.7, legend_y, 0.3, 0.25, GREEN)
add_text(slide, 1.1, legend_y - 0.02, 1.5, 0.3, '계획', size=11, color=GRAY)
add_sharp_rect(slide, 2.2, legend_y, 0.3, 0.25, ORANGE)
add_text(slide, 2.6, legend_y - 0.02, 1.5, 0.3, '실제 진행', size=11, color=GRAY)


# ────────────────────────────────────────────────
# Slide 9: 향후 보완 과제
# ────────────────────────────────────────────────
slide = blank_slide()
set_bg(slide, DARK)
slide_title(slide, '향후 보완 과제')

future_items = [
    ('OIDC 기반 SSO 통합 인증',
     '서비스별 자동 로그인 — Nextcloud / Paperless / ArchiveBox 통합'),
    ('모바일 반응형 최적화',
     '스마트폰·태블릿 환경에서 포탈 전체 기능 접근 가능'),
    ('자동 백업 및 장애 복구 체계',
     '정기 스냅샷 + 복구 절차 자동화'),
    ('모니터링 알림 시스템',
     '이상 감지 시 자동 통보 (Discord / 이메일 Webhook)'),
]
for i, (title, desc) in enumerate(future_items):
    y = 1.7 + i * 1.3
    add_rect(slide, 0.7, y, 11.9, 1.1, CARD)
    add_sharp_rect(slide, 0.7, y, 0.08, 1.1, BLUE)
    add_text(slide, 1.1, y + 0.1, 11.0, 0.4,
             title, size=18, color=WHITE, bold=True)
    add_text(slide, 1.1, y + 0.6, 11.0, 0.4,
             desc, size=14, color=GRAY)


# ────────────────────────────────────────────────
# Slide 10: 감사합니다
# ────────────────────────────────────────────────
slide = blank_slide()
set_bg(slide, DARK)

add_sharp_rect(slide, 0, 0, 0.5, 7.5, BLUE)
add_sharp_rect(slide, 12.833, 0, 0.5, 7.5, BLUE)

add_text(slide, 1, 2.3, 11.333, 1.2,
         '감사합니다',
         size=56, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
add_sharp_rect(slide, 4.5, 3.9, 4.333, 0.05, BLUE)
add_text(slide, 1, 4.2, 11.333, 0.6,
         '살려조  |  팀장 정성현',
         size=20, color=GRAY, align=PP_ALIGN.CENTER)


# ────────────────────────────────────────────────
# 저장
# ────────────────────────────────────────────────
out_path = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                        'CHEEZE_캡스톤디자인_기말발표_2026-06-11.pptx')
prs.save(out_path)
print(f'PPT saved: {out_path}')
